import calendar
import datetime
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ====================================================
# ESTILOS OPTIMIZADOS PARA IMPRESIÓN EN BLANCO Y NEGRO (B/N)
# ====================================================
FAMILIA_FUENTE = "Aptos"

FUENTE_TITULO = Font(name=FAMILIA_FUENTE, size=14, bold=True, color="000000")
FUENTE_HEADER = Font(name=FAMILIA_FUENTE, size=11, bold=True, color="000000")
FUENTE_SEMANA_ORD = Font(name=FAMILIA_FUENTE, size=8.5, bold=False, color="000000")
FUENTE_FECHA_SUB = Font(name=FAMILIA_FUENTE, size= 10, bold=True, color="000000")
FUENTE_SECCION = Font(name=FAMILIA_FUENTE, size=9.5, bold=True, color="000000")
FUENTE_ETIQUETA = Font(name=FAMILIA_FUENTE, size=10, bold=True, color="000000")
FUENTE_CELDA = Font(name=FAMILIA_FUENTE, size=10.5, color="000000")

FILL_ENCABEZADO = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
FILL_SECCION = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")

BORDE_NEGRO = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)

ALINEACION_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALINEACION_IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)


def aplicar_borde_rango(ws, r_ini: int, r_fin: int, c_ini: int, c_fin: int, borde: Border):
    """Asegura que todas las celdas de un bloque tengan sus cuatro bordes cerrados."""
    for r in range(r_ini, r_fin + 1):
        for c in range(c_ini, c_fin + 1):
            ws.cell(row=r, column=c).border = borde


def obtener_siguiente_grupo(ruta_estado: str = "data/estado.json") -> int:
    """Lee el grupo de limpieza que corresponde a la primera semana del siguiente mes."""
    archivo = Path(ruta_estado)
    if archivo.exists():
        try:
            with open(archivo, "r", encoding="utf-8") as f:
                datos = json.load(f)
                return datos.get("siguiente_grupo_limpieza", 1)
        except Exception:
            return 1
    return 1


def guardar_estado(anio: int, mes: int, ultimo_grupo: int, ruta_estado: str = "data/estado.json"):
    """Guarda en memoria el estado para la rotación automática de los 4 grupos."""
    archivo = Path(ruta_estado)
    archivo.parent.mkdir(parents=True, exist_ok=True)
    siguiente_grupo = (ultimo_grupo % 4) + 1
    datos = {
        "ultimo_anio": anio,
        "ultimo_mes": mes,
        "ultimo_grupo_limpieza": ultimo_grupo,
        "siguiente_grupo_limpieza": siguiente_grupo,
    }
    with open(archivo, "w", encoding="utf-8") as f:
        json.dump(datos, f, indent=4, ensure_ascii=False)


def calcular_semanas_mes(anio: int, mes: int) -> list[dict]:
    """Calcula las semanas de reunión del mes en función de los días de reunión (Martes a Domingo)."""
    num_dias = calendar.monthrange(anio, mes)[1]
    martes_mes = [
        datetime.date(anio, mes, d)
        for d in range(1, num_dias + 1)
        if datetime.date(anio, mes, d).weekday() == 1
    ]

    mes_corto = [
        "", "Ene", "Feb", "Mar", "Abr", "May", "Jun",
        "Jul", "Ago", "Sep", "Oct", "Nov", "Dic",
    ]
    ordinales = ["1ra . semana", "2da. semana", "3ra. semana", "4ta. semana", "5ta. semana"]

    semanas = []
    for idx, mar in enumerate(martes_mes):
        dom = mar + datetime.timedelta(days=5)

        if mar.month == dom.month:
            etiq_av = f"{mar.day:02d}-{dom.day:02d} {mes_corto[mar.month]}"
            sub_fecha = f"{mar.day:02d} - {dom.day:02d} {mes_corto[dom.month]}."
        else:
            etiq_av = f"{mar.day:02d} {mes_corto[mar.month]} - {dom.day:02d} {mes_corto[dom.month]}"
            sub_fecha = f"{mar.day:02d} {mes_corto[mar.month]}. - {dom.day:02d} {mes_corto[dom.month]}."

        semanas.append({
            "semana": idx + 1,
            "ordinal": ordinales[idx],
            "fecha_semana": sub_fecha,
            "audio_video": etiq_av,
        })

    return semanas


def exportar_programa_excel(
    mes_asignaciones: dict[int, dict],
    anio: int,
    mes: int,
    grupo_inicio_limpieza: int | None = None,
    ruta_salida: str = "salida/programa_mes.xlsx",
    ruta_estado: str = "data/estado.json",
) -> dict:
    """
    Genera el archivo Excel optimizado:
    Hoja 1: 'Programa General' en una tabla matriz única para el tablero (A4 horizontal).
    Hoja 2: 'Audio y Video' para el equipo técnico digital.
    """
    if grupo_inicio_limpieza is None:
        grupo_inicio = obtener_siguiente_grupo(ruta_estado)
    else:
        grupo_inicio = grupo_inicio_limpieza

    wb = openpyxl.Workbook()
    semanas = calcular_semanas_mes(anio, mes)
    num_semanas = len(semanas)

    nombres_meses = [
        "", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
        "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
    ]
    nombre_mes_str = nombres_meses[mes]

    # ====================================================
    # HOJA 1: PROGRAMA GENERAL (Matriz única para el tablero)
    # ====================================================
    ws_gen = wb.active
    ws_gen.title = "Programa General"
    ws_gen.views.sheetView[0].showGridLines = True

    # Configuración de página A4 Horizontal ajustada a 1 sola hoja
    ws_gen.page_setup.orientation = ws_gen.ORIENTATION_LANDSCAPE
    ws_gen.page_setup.paperSize = ws_gen.PAPERSIZE_A4
    ws_gen.page_setup.fitToWidth = 1
    ws_gen.page_setup.fitToHeight = 1
    ws_gen.sheet_properties.pageSetUpPr.fitToPage = True

    # Margen izquierdo y superior para recortar solo dos lados
    ws_gen.page_margins.left = 0.35
    ws_gen.page_margins.top = 0.35
    ws_gen.page_margins.right = 0.4
    ws_gen.page_margins.bottom = 0.4

    ws_gen.print_options.horizontalCentered = False
    ws_gen.print_options.verticalCentered = False

    # 1. TÍTULO PRINCIPAL
    ws_gen.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_semanas + 1)
    c_tit = ws_gen.cell(
        row=1,
        column=1,
        value=f"PROGRAMA DE ASIGNACIONES - {nombre_mes_str} {anio}",
    )
    c_tit.font = FUENTE_TITULO
    c_tit.alignment = ALINEACION_CENTRO
    ws_gen.row_dimensions[1].height = 35

    # 2. ENCABEZADOS DE SEMANA (Filas 2 y 3)
    ws_gen.merge_cells("A2:A3")
    c_asig = ws_gen.cell(row=2, column=1, value="ASIGNACIÓN")
    c_asig.font = FUENTE_HEADER
    c_asig.alignment = ALINEACION_CENTRO
    c_asig.fill = FILL_ENCABEZADO
    aplicar_borde_rango(ws_gen, 2, 3, 1, 1, BORDE_NEGRO)

    for idx, s in enumerate(semanas, start=2):
        c_ord = ws_gen.cell(row=2, column=idx, value=s["ordinal"])
        c_ord.font = FUENTE_SEMANA_ORD
        c_ord.alignment = ALINEACION_CENTRO
        c_ord.fill = FILL_ENCABEZADO
        c_ord.border = BORDE_NEGRO

        c_fec = ws_gen.cell(row=3, column=idx, value=s["fecha_semana"])
        c_fec.font = FUENTE_FECHA_SUB
        c_fec.alignment = ALINEACION_CENTRO
        c_fec.fill = FILL_ENCABEZADO
        c_fec.border = BORDE_NEGRO

    ws_gen.row_dimensions[2].height = 18
    ws_gen.row_dimensions[3].height = 25

    # 3. FILAS Y SECCIONES UNIFICADAS
    estructura_tabla = [

        # SECCIÓN 1: AUDITORIO Y SALA
        (True, "AUDITORIO Y PLATAFORMA", None, 24),
        (False, "Acomodadores", "acomodador", 45),
        (False, "Plataforma", "plataforma", 30),
        (False, "Micrófonos", "microfonos", 45),
        (False, "Limpieza del Salón", "limpieza", 30),
        (False, "Hospitalidad", "hospitalidad", 30),

        # SECCIÓN 2: PLATAFORMA Y LECTURAS
        (True, "PRESIDENCIA Y LECTURAS", None, 24),
        (False, "Lector Estudio Bíblico", "lector_martes", 30),
        (False, "Presidente", "presidente", 30),
        (False, "Lector de La Atalaya", "lector_domingo", 30),

    ]

    r_actual = 4
    for es_sec, label, key, alto_fila in estructura_tabla:
        if es_sec:
            ws_gen.merge_cells(start_row=r_actual, start_column=1, end_row=r_actual, end_column=num_semanas + 1)
            for c_idx in range(1, num_semanas + 2):
                c = ws_gen.cell(row=r_actual, column=c_idx)
                c.fill = FILL_SECCION
                c.border = BORDE_NEGRO
            c_sec = ws_gen.cell(row=r_actual, column=1, value=label)
            c_sec.font = FUENTE_SECCION
            c_sec.alignment = ALINEACION_CENTRO
        else:
            c_lbl = ws_gen.cell(row=r_actual, column=1, value=label)
            c_lbl.font = FUENTE_ETIQUETA
            c_lbl.alignment = ALINEACION_IZQ
            c_lbl.border = BORDE_NEGRO

            for s_idx in range(1, num_semanas + 1):
                if key == "limpieza":
                    num_limpieza = ((grupo_inicio - 1 + (s_idx - 1)) % 4) + 1
                    val = f"Grupo # {num_limpieza}"
                elif key == "hospitalidad":
                    num_limpieza = ((grupo_inicio - 1 + (s_idx - 1)) % 4) + 1
                    num_hosp = ((num_limpieza - 2) % 4) + 1
                    val = f"Grupo # {num_hosp}"
                else:
                    nombres = mes_asignaciones.get(s_idx, {}).get(key, [])
                    val = "\n".join(nombres) if nombres else "—"

                c = ws_gen.cell(row=r_actual, column=s_idx + 1, value=val)
                c.font = FUENTE_CELDA
                c.alignment = ALINEACION_CENTRO
                c.border = BORDE_NEGRO

        ws_gen.row_dimensions[r_actual].height = alto_fila
        r_actual += 1

    # Anchos equilibrados
    ws_gen.column_dimensions["A"].width = 21.5
    for c_idx in range(2, num_semanas + 2):
        ws_gen.column_dimensions[get_column_letter(c_idx)].width = 19.5

    # ====================================================
    # HOJA 2: AUDIO Y VIDEO (Solo digital)
    # ====================================================
    ws_av = wb.create_sheet(title="Audio y Video")
    ws_av.views.sheetView[0].showGridLines = True

    ws_av.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_semanas + 1)
    c_tit_av = ws_av.cell(
        row=1,
        column=1,
        value=f"PROGRAMA DE ASIGNACIONES AUDIO Y VIDEO - {nombre_mes_str} {anio}",
    )
    c_tit_av.font = FUENTE_TITULO
    c_tit_av.alignment = ALINEACION_CENTRO
    ws_av.row_dimensions[1].height = 28

    c_fec_av = ws_av.cell(row=2, column=1, value="Fecha")
    c_fec_av.font = FUENTE_HEADER
    c_fec_av.alignment = ALINEACION_CENTRO
    c_fec_av.fill = FILL_ENCABEZADO
    c_fec_av.border = BORDE_NEGRO

    for idx, s in enumerate(semanas, start=2):
        c = ws_av.cell(row=2, column=idx, value=s["audio_video"])
        c.font = FUENTE_HEADER
        c.alignment = ALINEACION_CENTRO
        c.fill = FILL_ENCABEZADO
        c.border = BORDE_NEGRO
    ws_av.row_dimensions[2].height = 26

    for r_idx, (label, key) in enumerate([("Audio", "audio"), ("Video", "video")], start=3):
        c = ws_av.cell(row=r_idx, column=1, value=label)
        c.font = FUENTE_ETIQUETA
        c.alignment = ALINEACION_IZQ
        c.border = BORDE_NEGRO

        for s_idx in range(1, num_semanas + 1):
            nombres = mes_asignaciones.get(s_idx, {}).get(key, [])
            val = "\n".join(nombres) if nombres else "—"
            cn = ws_av.cell(row=r_idx, column=s_idx + 1, value=val)
            cn.font = FUENTE_CELDA
            cn.alignment = ALINEACION_CENTRO
            cn.border = BORDE_NEGRO

        ws_av.row_dimensions[r_idx].height = 26

    ws_av.column_dimensions["A"].width = 16
    for idx in range(2, num_semanas + 2):
        ws_av.column_dimensions[get_column_letter(idx)].width = 20

    # Guardar archivo Excel
    ruta = Path(ruta_salida)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)

    # Actualizar estado para el mes siguiente
    ultimo_grupo = ((grupo_inicio - 1 + (num_semanas - 1)) % 4) + 1
    guardar_estado(anio, mes, ultimo_grupo, ruta_estado)

    col_fin = get_column_letter(num_semanas + 1)
    fila_fin = r_actual - 1

    return {
        "ruta_excel": str(ruta.resolve()),
        "num_semanas": num_semanas,
        "rango_tabla": f"A1:{col_fin}{fila_fin}",
        "hoja": "Programa General"
    }
